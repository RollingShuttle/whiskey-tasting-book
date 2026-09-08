"""
test_master_write.py — the surgical append, in isolation.

    python test_master_write.py

Builds a small workbook shaped like the real master — a table defined over a range far larger than
its data, ABV as a filled-down formula — and injects fake `xl/media/` and `xl/richData/` parts plus
a calcChain, so that preserving them is something the tests can actually prove.

The real `Whiskey Collection.xlsx` is never opened by this module or its tests.
"""
import hashlib
import shutil
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

import collection as collection_mod
import master_write as mw

HEADERS = ["Bottle Code", "Distillery", "Name", "Type", "Region", "Notes", "Release Year", "Age",
           "Rarity", "Proof", "ABV", "Entry Proof", "Proof Diff", "Conc. Ratio", "Size (ml)",
           "Status", "Paid"]
ROWS = [
    ("B-1", "Alpha Co", "One", "Bourbon", "America", None, 2024, 5, "Common",
     100.0, None, 110.0, None, None, 750.0, "Opened", 50.0),
    ("B-2", "Beta Co", "Two", "Rye", "America", None, 2023, 7, "Rare",
     110.0, None, 120.0, None, None, 750.0, "Unopened", 90.0),
]
MEDIA = {"xl/media/image1.png": b"\x89PNG\r\n\x1a\n" + b"photo-one" * 40,
         "xl/media/image2.png": b"\x89PNG\r\n\x1a\n" + b"photo-two" * 55}
RICH = {"xl/richData/rdrichvalue.xml": b'<?xml version="1.0"?><rvData count="2"/>',
        "xl/richData/rdrichvaluestructure.xml": b'<?xml version="1.0"?><rvStructures count="1"/>'}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build_fixture(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Bottle"
    for i, h in enumerate(HEADERS, start=2):            # columns B..R
        ws.cell(row=8, column=i, value=h)
    for r, row in enumerate(ROWS, start=9):
        for i, v in enumerate(row, start=2):
            ws.cell(row=r, column=i, value=v)
        # ABV is a plain filled-down formula with an A1 reference, exactly like the real sheet
        ws.cell(row=r, column=12, value=f'=IF(K{r}="","",K{r}/2)')
    ws.add_table(Table(displayName="Collection", ref="B8:R308"))
    wb.save(path)
    wb.close()

    # inject the parts that make this workbook dangerous to round-trip
    with zipfile.ZipFile(path) as z:
        entries = {i.filename: z.read(i.filename) for i in z.infolist()}
    entries.update(MEDIA)
    entries.update(RICH)
    entries["xl/calcChain.xml"] = b'<?xml version="1.0"?><calcChain/>'
    entries["[Content_Types].xml"] = entries["[Content_Types].xml"].decode().replace(
        "</Types>", '<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.'
                    'openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>'
    ).encode()
    entries["xl/_rels/workbook.xml.rels"] = entries["xl/_rels/workbook.xml.rels"].decode().replace(
        "</Relationships>", '<Relationship Id="rIdCalc" Type="http://schemas.openxmlformats.org/'
                            'officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>'
                            "</Relationships>"
    ).encode()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for name, data in entries.items():
            out.writestr(name, data)


class WriteCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="mw-test-"))
        self.book = self.tmp / "Whiskey Collection.xlsx"
        self.backups = self.tmp / "backups"
        build_fixture(self.book)

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def coll(self):
        return collection_mod.Collection(self.book).load()

    def append(self, **fields):
        return mw.append_row(self.book, "Bottle", fields, backup_dir=self.backups,
                             coll=self.coll())

    def parts(self, path=None):
        with zipfile.ZipFile(path or self.book) as z:
            return {i.filename: z.read(i.filename) for i in z.infolist()}


class TestPureHelpers(unittest.TestCase):
    def test_column_letters_round_trip(self):
        for n, letters in [(1, "A"), (2, "B"), (18, "R"), (26, "Z"), (27, "AA"), (28, "AB")]:
            self.assertEqual(mw.index_to_col(n), letters)
            self.assertEqual(mw.col_to_index(letters), n)

    def test_a1_references_shift_but_structured_ones_do_not(self):
        self.assertEqual(mw.shift_formula('=IF(K9="","",K9/2)', 9, 153),
                         '=IF(K153="","",K153/2)')
        structured = "=Collection[[#This Row],[Proof]]-Collection[[#This Row],[Entry Proof]]"
        self.assertEqual(mw.shift_formula(structured, 9, 153), structured)

    def test_shifting_does_not_maul_a_lookalike(self):
        """Row 9 must not match inside 19, 90 or a defined name."""
        self.assertEqual(mw.shift_formula("=K19+K90", 9, 153), "=K19+K90")

    def test_a_formula_cell_carries_no_cached_value(self):
        xml = mw.formula_cell("L153", '=IF(K153="","",K153/2)')
        self.assertIn("<f>", xml)
        self.assertNotIn("<v>", xml)

    def test_strings_go_in_inline(self):
        self.assertIn('t="inlineStr"', mw.inline_string_cell("C153", "Willett"))
        self.assertIn("&amp;", mw.inline_string_cell("C153", "Smith & Sons"))

    def test_cells_are_written_in_ascending_column_order(self):
        xml = '<sheetData><row r="5"><c r="B5"><v>1</v></c></row></sheetData>'
        out = mw.replace_row(xml, 5, {"R": '<c r="R5"><v>9</v></c>',
                                      "C": '<c r="C5"><v>2</v></c>'})
        self.assertLess(out.index('r="B5"'), out.index('r="C5"'))
        self.assertLess(out.index('r="C5"'), out.index('r="R5"'))

    def test_an_untouched_cell_in_the_row_survives(self):
        xml = '<sheetData><row r="5"><c r="B5" s="3"/></row></sheetData>'
        out = mw.replace_row(xml, 5, {"C": '<c r="C5"><v>2</v></c>'})
        self.assertIn('r="B5"', out)

    def test_a_missing_row_is_created_in_ascending_position(self):
        xml = '<sheetData><row r="4"/><row r="9"/></sheetData>'
        out = mw.replace_row(xml, 5, {"B": '<c r="B5"><v>1</v></c>'})
        self.assertLess(out.index('r="4"'), out.index('r="5"'))
        self.assertLess(out.index('r="5"'), out.index('r="9"'))


class TestAppend(WriteCase):
    def test_the_new_row_lands_with_the_next_code(self):
        result = self.append(Distillery="Gamma Co", Name="Three", Type="Rum", Proof=101.4)
        self.assertEqual(result["code"], "B-3")
        self.assertEqual(result["row"], 11)          # first empty row inside the table
        self.assertEqual((result["rows_before"], result["rows_after"]), (2, 3))

        rows = {r["code"]: r for r in self.coll().rows["Bottle"]}
        self.assertEqual(set(rows), {"B-1", "B-2", "B-3"})
        self.assertEqual(rows["B-3"]["distillery"], "Gamma Co")
        self.assertEqual(rows["B-3"]["proof"], 101.4)

    def test_every_photo_and_rich_value_survives_byte_for_byte(self):
        """This is the whole reason the module exists."""
        self.append(Distillery="Gamma Co", Name="Three")
        after = self.parts()
        for name, blob in {**MEDIA, **RICH}.items():
            self.assertIn(name, after, f"{name} was lost")
            self.assertEqual(after[name], blob, f"{name} changed")

    def test_shared_strings_are_never_touched(self):
        before = self.parts().get("xl/sharedStrings.xml")
        self.append(Distillery="Gamma Co", Name="Three")
        self.assertEqual(self.parts().get("xl/sharedStrings.xml"), before)

    def test_the_table_ref_does_not_move(self):
        before = self.coll().tables["Bottle"].ref
        self.append(Distillery="Gamma Co", Name="Three")
        self.assertEqual(self.coll().tables["Bottle"].ref, before)

    def test_calc_chain_is_dropped_with_its_override_and_relationship(self):
        self.append(Distillery="Gamma Co", Name="Three")
        after = self.parts()
        self.assertNotIn("xl/calcChain.xml", after)
        self.assertNotIn(b"calcChain.xml", after["[Content_Types].xml"])
        self.assertNotIn(b"calcChain.xml", after["xl/_rels/workbook.xml.rels"])

    def test_full_calc_on_load_is_set(self):
        self.append(Distillery="Gamma Co", Name="Three")
        self.assertIn(b'fullCalcOnLoad="1"', self.parts()["xl/workbook.xml"])

    def test_abv_is_written_as_a_formula_with_the_row_shifted(self):
        """Bottle's ABV is filled-down, not a declared calculated column, so Excel will not add
        it for us — the writer has to copy it and move the row reference (SPEC.md §8.3)."""
        self.append(Distillery="Gamma Co", Name="Three", Proof=101.4)
        sheet = self.parts()["xl/worksheets/sheet1.xml"].decode()
        row = sheet[sheet.index('<row r="11"'):]
        row = row[:row.index("</row>")]
        self.assertIn("<f>", row)
        self.assertIn("K11", row)
        self.assertNotIn("K9", row)

    def test_two_appends_take_consecutive_codes(self):
        self.append(Distillery="Gamma Co", Name="Three")
        second = self.append(Distillery="Delta Co", Name="Four")
        self.assertEqual(second["code"], "B-4")
        self.assertEqual(second["row"], 12)
        self.assertEqual(len(self.coll().rows["Bottle"]), 4)

    def test_a_backup_is_taken_before_the_write(self):
        before = sha(self.book)
        result = self.append(Distillery="Gamma Co", Name="Three")
        self.assertEqual(sha(result["backup"]), before, "the backup is the pre-write file")


class TestGuards(WriteCase):
    def test_an_encounter_is_never_written_to_the_collection(self):
        """A bar pour is not inventory (SPEC.md §2.1, §8.2)."""
        before = sha(self.book)
        with self.assertRaises(mw.MasterWriteError) as ctx:
            self.append(**{"Bottle Code": "X-4", "Name": "Bar pour"})
        self.assertIn("X-", str(ctx.exception))
        self.assertEqual(sha(self.book), before, "nothing may be written")

    def test_it_refuses_while_excel_holds_the_file(self):
        lock = self.book.with_name(f"~${self.book.name}")
        lock.write_bytes(b"")
        before = sha(self.book)
        with self.assertRaises(mw.MasterWriteError) as ctx:
            self.append(Distillery="Gamma Co", Name="Three")
        self.assertIn("open in Excel", str(ctx.exception))
        self.assertEqual(sha(self.book), before)

    def test_a_full_table_is_refused_rather_than_grown(self):
        coll = self.coll()
        coll.tables["Bottle"].ref = "B8:R10"          # only rows 9 and 10 — both filled
        before = sha(self.book)
        with self.assertRaises(mw.MasterWriteError) as ctx:
            mw.append_row(self.book, "Bottle", {"Name": "Three"},
                          backup_dir=self.backups, coll=coll)
        self.assertIn("full", str(ctx.exception))
        self.assertEqual(sha(self.book), before)

    def test_a_failed_verification_leaves_the_master_untouched(self):
        before = sha(self.book)
        original = mw._verify
        mw._verify = lambda *a, **k: (_ for _ in ()).throw(AssertionError("boom"))
        try:
            with self.assertRaises(mw.MasterWriteError):
                self.append(Distillery="Gamma Co", Name="Three")
        finally:
            mw._verify = original
        self.assertEqual(sha(self.book), before, "a failed write must not change the master")
        self.assertEqual(list(self.tmp.glob(".tmp-*")), [], "no debris left behind")


if __name__ == "__main__":
    unittest.main(verbosity=2)
