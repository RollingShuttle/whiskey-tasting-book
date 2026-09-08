"""
quickentry.py — drain the `Quick Entry` sheet of the rollup workbook.

The phone lane (SPEC.md §1.2, §9.1 fallback). A row is typed one-handed, in Excel for iOS, into
the generated workbook: `date · display_name · barrel_id · nose · palate · finish · notes`. No
scores — the point is to capture the pour before it is forgotten and score it properly later.

On the PC, draining a row turns it into a **draft** tasting in the journal, with the free text
mapped onto the matching category notes. A draft carries no scores, so it can never move a career
score; it shows up as something waiting to be scored.

Two rules from the spec that shape this module:

  * **Never match on a name that is not unique** (§2). `Distillery + Name` collides all over the
    collection, so an ambiguous row is refused rather than guessed at. Typing the bottle code is
    always unambiguous and always wins.
  * **Rows that fail to match are left in place and flagged**, never dropped (§1.2). drain()
    returns them so the caller can write them back into the regenerated sheet with the reason.

The rollup workbook is *our own* generated file — no images, no rich values — so opening it with
openpyxl is safe. This module never goes near `Whiskey Collection.xlsx` (§1.1).
"""
from __future__ import annotations

import re
from datetime import date as _date, datetime
from pathlib import Path

import openpyxl

SHEET = "Quick Entry"
HEADERS = ["date", "display_name", "barrel_id", "nose", "palate", "finish", "notes"]
CODE_RE = re.compile(r"^[BMSX]-\d+$", re.IGNORECASE)

# Quick Entry free text -> the category whose notes it belongs to.
NOTE_MAP = {"nose": "aroma", "palate": "flavor", "finish": "finish", "notes": "overall"}


def _txt(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, _date):
        return v.strftime("%Y-%m-%d")
    s = str(v).strip()
    return s or None


def _norm(v):
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


def read_rows(workbook_path):
    """Every non-empty row a human typed into the Quick Entry sheet."""
    path = Path(workbook_path)
    if not path.exists():
        return []
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            return []
        rows = [list(r) for r in wb[SHEET].iter_rows(values_only=True)]
    finally:
        wb.close()
    if not rows:
        return []

    head = [_norm(c) for c in rows[0]]
    out = []
    for n, raw in enumerate(rows[1:], start=2):
        rec = {}
        for i, name in enumerate(head):
            if name and i < len(raw):
                rec[name] = _txt(raw[i])
        # `problem` is read back so the UI can say why a row was refused last time. It is display
        # only — drain() re-derives it, and it never reaches the journal.
        if not any(rec.values()):
            continue                      # a blank row waiting to be typed into
        rec["_row"] = n
        out.append(rec)
    return out


def match_spirit(text, spirits):
    """(code, problem). Exactly one of the two is set.

    A bottle code wins outright. Otherwise the name has to identify one spirit and only one —
    guessing between two bottles that share a name is how a score ends up on the wrong bottle.
    """
    q = _norm(text)
    if not q:
        return None, "no name typed"

    if CODE_RE.match(q):
        for s in spirits:
            if _norm(s.get("code")) == q:
                return s["code"], None
        return None, f"no spirit with code {text.strip()}"

    exact = [s for s in spirits
             if _norm(s.get("display_name")) == q or _norm(s.get("name")) == q]
    if len(exact) == 1:
        return exact[0]["code"], None
    if len(exact) > 1:
        return None, (f"ambiguous — {len(exact)} spirits share that name "
                      f"({', '.join(s['code'] for s in exact[:4])}…). Type the bottle code.")

    partial = [s for s in spirits if q in _norm(s.get("display_name"))]
    if len(partial) == 1:
        return partial[0]["code"], None
    if len(partial) > 1:
        return None, (f"ambiguous — {len(partial)} spirits match "
                      f"({', '.join(s['code'] for s in partial[:4])}…). Type the bottle code.")
    return None, "no match in the collection"


def notes_from(row):
    """Map the phone's three impressions onto category notes; `notes` is the overall note."""
    out = {}
    for src, key in NOTE_MAP.items():
        if row.get(src):
            out[key] = row[src]
    return out


def drain(workbook_path, journal, spirits):
    """Turn every matchable Quick Entry row into a draft tasting.

    Returns {"drained": [...], "unmatched": [...], "counts": {...}}. The caller regenerates the
    workbook with `unmatched` so the flagged rows stay on the sheet to be corrected.
    """
    rows = read_rows(workbook_path)
    drained, unmatched = [], []

    for row in rows:
        code, problem = match_spirit(row.get("display_name"), spirits)
        if problem:
            unmatched.append(dict(row, problem=problem))
            continue
        rec = journal.write_tasting(
            spirit_id=code,
            scores=None,                       # a Quick Entry row has no numbers yet
            notes=notes_from(row),
            date=row.get("date"),
            barrel_id=row.get("barrel_id"),
            status="draft",
            entered_from="quick-entry",
            include_in_average=False,
        )
        drained.append({"tasting_id": rec["tasting_id"], "spirit_id": code,
                        "display_name": row.get("display_name"), "row": row.get("_row")})

    return {
        "drained": drained,
        "unmatched": unmatched,
        "counts": {"read": len(rows), "drained": len(drained), "unmatched": len(unmatched)},
    }
