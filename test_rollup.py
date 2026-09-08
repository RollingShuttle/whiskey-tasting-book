"""
test_rollup.py — the generated workbook, in isolation. Temp dirs only; the master collection
workbook is never opened by this module or its tests.

    python test_rollup.py
"""
import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

import rollup
import rubric as rubric_mod
import store
from test_rubric import EXAMPLE_CARD, EXAMPLE_SITTINGS


class RollupCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp(prefix="rollup-test-"))
        self.j = store.Journal(self.tmp / "journal", rubric_mod.load_rubric()).ensure()
        self.out = self.tmp / "Whiskey Tastings.xlsx"
        self.backups = self.tmp / "backups"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def build(self):
        return rollup.build(self.j, self.out, self.backups)

    def load(self):
        return openpyxl.load_workbook(self.out, data_only=True)

    def seed_sittings(self):
        for s in EXAMPLE_SITTINGS:
            self.j.write_tasting(spirit_id="B-18", scores=s["scores"], date=s["date"],
                                 include_in_average=s["include_in_average"],
                                 notes={"aroma": "caramel, dark cherry"})


class TestStructure(RollupCase):
    def test_empty_journal_still_produces_a_valid_workbook(self):
        self.build()
        wb = self.load()
        self.assertEqual(wb.sheetnames,
                         ["Tastings", "Quick Entry", "Careers", "Sessions", "Encounters",
                          "Pending", "About"])
        wb.close()

    def test_every_rubric_category_gets_a_score_and_a_notes_column(self):
        self.build()
        wb = self.load()
        head = [c.value for c in wb["Tastings"][1]]
        for c in self.j.rubric.categories:
            self.assertIn(c.label, head, f"{c.label} score column")
            self.assertIn(f"{c.label} notes", head, f"{c.label} notes column")
        self.assertIn("total", head)
        self.assertIn("medal", head)
        self.assertIn("counted", head)
        wb.close()

    def test_about_sheet_warns_that_the_file_is_generated(self):
        self.build()
        wb = self.load()
        text = " ".join(str(c.value) for row in wb["About"].iter_rows() for c in row)
        self.assertIn("do not hand-edit", text)
        self.assertIn("roundings drift", text)          # the §3.6 trap, stated in the file itself
        wb.close()


class TestContent(RollupCase):
    def test_sittings_land_with_totals_medals_and_notes(self):
        self.seed_sittings()
        summary = self.build()
        self.assertEqual(summary["tastings"], 4)

        wb = self.load()
        ws = wb["Tastings"]
        head = [c.value for c in ws[1]]
        rows = [dict(zip(head, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 4)
        self.assertEqual({r["spirit_id"] for r in rows}, {"B-18"})
        self.assertEqual(sorted(r["total"] for r in rows), [76, 87, 89, 91])
        self.assertEqual({r["medal"] for r in rows}, {"Silver", "Gold", "Diamond"})
        self.assertEqual(sorted(r["counted"] for r in rows), ["no", "yes", "yes", "yes"])
        self.assertEqual(rows[0]["Aroma notes"], "caramel, dark cherry")
        self.assertEqual(rows[0]["Flavor"], 18)          # newest first
        wb.close()

    def test_careers_sheet_shows_the_aggregate_not_the_last_sitting(self):
        self.seed_sittings()
        self.build()
        wb = self.load()
        ws = wb["Careers"]
        head = [c.value for c in ws[1]]
        row = dict(zip(head, [c.value for c in ws[2]]))
        self.assertEqual(row["spirit_id"], "B-18")
        self.assertEqual(row["career_score"], 89.0)
        self.assertEqual(row["medal"], "Gold")
        self.assertEqual(row["sittings_counted"], 3)
        self.assertEqual(row["sittings_total"], 4)
        self.assertEqual((row["best"], row["worst"]), (91, 87))
        self.assertEqual(row["Aroma mean"], 8.7)
        self.assertEqual(row["Aroma range"], "8–9")
        self.assertEqual(row["Body range"], "9")          # never varied
        wb.close()

    def test_a_tombstoned_sitting_leaves_the_workbook(self):
        self.seed_sittings()
        victim = self.j.tastings()[0]["tasting_id"]
        self.j.delete_tasting(victim, reason="test")
        self.assertEqual(self.build()["tastings"], 3)

    def test_encounters_and_pending_appear(self):
        self.j.write_encounter(name="Example Cask 10 Year", venue="Bar")
        self.j.assign_encounter_codes()
        self.j.write_pending_bottle(sheet="Bottle", fields={"Distillery": "Example Distillery"})
        s = self.build()
        self.assertEqual((s["encounters"], s["pending"]), (1, 1))
        wb = self.load()
        self.assertEqual(wb["Encounters"]["A2"].value, "X-1")
        self.assertIn("Example Distillery", wb["Pending"]["E2"].value)
        wb.close()


    def test_sessions_sheet_lists_flights_with_their_pour_counts(self):
        sid = self.j.write_session(title="Thursday flight", location="Home",
                                   blind=True)["session_id"]
        for pos, code in enumerate(["B-1", "B-2"], start=1):
            self.j.write_tasting(spirit_id=code, scores=EXAMPLE_CARD,
                                 session_id=sid, flight_pos=pos)
        summary = self.build()
        self.assertEqual(summary["sessions"], 1)

        wb = self.load()
        ws = wb["Sessions"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        self.assertEqual(row["session_id"], sid)
        self.assertEqual(row["title"], "Thursday flight")
        self.assertEqual(row["blind"], "yes")
        self.assertEqual(row["pours"], 2)
        wb.close()

    def test_a_flight_links_its_pours_by_session_id(self):
        sid = self.j.write_session(title="Linked")["session_id"]
        self.j.write_tasting(spirit_id="B-1", scores=EXAMPLE_CARD, session_id=sid, flight_pos=1)
        self.build()
        wb = self.load()
        ws = wb["Tastings"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        self.assertEqual(row["session_id"], sid)
        self.assertEqual(row["flight_pos"], 1)
        wb.close()


    def test_quick_entry_starts_empty_and_ready_to_type_into(self):
        summary = self.build()
        self.assertEqual(summary["quick_entry"], 0)
        wb = self.load()
        ws = wb["Quick Entry"]
        self.assertEqual([c.value for c in ws[1]][:4],
                         ["date", "display_name", "barrel_id", "nose"])
        self.assertEqual(ws.max_row, 1, "headers only until something is flagged")
        wb.close()

    def test_unmatched_quick_entry_rows_come_back_onto_the_sheet(self):
        """Rows that fail to match are left in place and flagged, never dropped (SPEC.md §1.2)."""
        rollup.build(self.j, self.out, self.backups, quick_rows=[
            {"date": "2026-09-08", "display_name": "Mystery pour", "nose": "smoke",
             "problem": "no match in the collection"}])
        wb = self.load()
        ws = wb["Quick Entry"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        self.assertEqual(row["display_name"], "Mystery pour")
        self.assertEqual(row["nose"], "smoke")
        self.assertEqual(row["problem"], "no match in the collection")
        wb.close()

    def test_a_draft_lands_unscored_and_uncounted(self):
        """A drained Quick Entry row has notes but no numbers; it must not disturb any career."""
        self.j.write_tasting(spirit_id="B-1", status="draft", entered_from="quick-entry",
                             notes={"aroma": "smoke"}, barrel_id="F664")
        self.build()
        wb = self.load()
        ws = wb["Tastings"]
        row = dict(zip([c.value for c in ws[1]], [c.value for c in ws[2]]))
        self.assertIsNone(row["total"])
        self.assertIsNone(row["Aroma"])
        self.assertEqual(row["counted"], "no")
        self.assertEqual(row["status"], "draft")
        self.assertEqual(row["Aroma notes"], "smoke")
        self.assertEqual(row["barrel_id"], "F664")
        wb.close()


class TestSafety(RollupCase):
    def test_refuses_to_write_while_excel_has_the_file_open(self):
        self.seed_sittings()
        self.build()
        lock = self.out.with_name(f"~${self.out.name}")
        lock.write_bytes(b"")
        with self.assertRaises(RuntimeError) as ctx:
            self.build()
        self.assertIn("open in Excel", str(ctx.exception))

    def test_backup_is_taken_before_each_overwrite_and_pruned(self):
        self.seed_sittings()
        self.build()                                     # nothing to back up yet
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 0)
        for _ in range(3):
            self.build()
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 3)

        rollup.build(self.j, self.out, self.backups, keep=2)
        self.assertEqual(len(list(self.backups.glob("*.xlsx"))), 2)

    def test_no_temp_debris_left_behind(self):
        self.seed_sittings()
        self.build()
        self.assertEqual(list(self.out.parent.glob(".tmp-*")), [])

    def test_regeneration_is_stable(self):
        """Deleting the workbook must be harmless — it rebuilds identically in content."""
        self.seed_sittings()
        first = self.build()
        self.out.unlink()
        second = self.build()
        self.assertEqual(first["tastings"], second["tastings"])
        self.assertEqual(first["careers"], second["careers"])
        wb = self.load()
        self.assertEqual(wb["Careers"]["B2"].value, 89.0)
        wb.close()

    def test_journal_is_not_modified_by_generating_the_rollup(self):
        self.seed_sittings()
        before = {p.name: p.read_bytes() for p in (self.j.root / "tastings").glob("*.json")}
        self.build()
        after = {p.name: p.read_bytes() for p in (self.j.root / "tastings").glob("*.json")}
        self.assertEqual(before, after, "the rollup must never write back to the journal")


if __name__ == "__main__":
    unittest.main(verbosity=1)
